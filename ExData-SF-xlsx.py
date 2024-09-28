import openpyxl

# 创建一个新的 Excel 工作簿和工作表
workbook = openpyxl.Workbook()
worksheet = workbook.active
worksheet.title = "SF-Numbers"

# 指定表头文本
headers = ["x_value", "NTrans", "Finesse", "Odess", "Palantir", "BiSearch"]

# 写入表头
for col_num, header in enumerate(headers, start=1):
    worksheet.cell(row=1, column=col_num, value=header)

# 指定文件名列表和后缀
suffix = "_WEB.txt"
file_names = ["Ntransform", "Finesse", "Odess", "Palantir", "BiSearchMultiTar"]

# 初始化列号，从第二列开始
col_num = 2

# 运行 5 次循环来填充后五列
max_rows = 0
for base_name in file_names:
    file_name = base_name + suffix
    row_num = 2
    with open(file_name, 'r') as file:
        for line in file:
            if line.startswith("SF number:"):
                # 提取第三列数据并转换为数字
                sf_number_str = line.split()[2]
                sf_number = float(sf_number_str) if '.' in sf_number_str else int(sf_number_str)
                # 写入 Excel 文件的指定列
                cell = worksheet.cell(row=row_num, column=col_num, value=sf_number)
                # 设置单元格格式为数字
                cell.number_format = '0.00' if isinstance(sf_number, float) else '0'
                row_num += 1
    max_rows = max(max_rows, row_num - 1)
    col_num += 1

# 填充第一列 x_value
for row_num in range(2, max_rows + 1):
    worksheet.cell(row=row_num, column=1, value=row_num - 1)

# 保存 Excel 文件
workbook.save("output.xlsx")

print("Excel 文件已生成：output.xlsx")